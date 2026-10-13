import math
import msvcrt
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent


def to_float(value, default=None):
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return default
    return number if math.isfinite(number) else default


def clock(seconds):
    h, rem = divmod(int(seconds), 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def ask_number(prompt, high):
    """Ask for an int between 0 and high."""
    while True:
        try:
            n = int(input(prompt))
        except ValueError:
            print("Enter a number.")
            continue
        if 0 <= n <= high:
            return n
        print(f"Enter a number from 0 to {high}.")


def choose_workbook():
    files = sorted(
        (f for f in BASE_DIR.glob("*.xlsx") if not f.name.startswith("~$")),
        key=lambda f: f.stat().st_mtime,
        reverse=True,
    )
    if not files:
        print("No Excel files found in this folder.")
        return None

    answer = input(f"Use '{files[0].name}'? (Y/n): ").strip().lower()
    if answer in ("", "y", "yes"):
        return files[0]
    if len(files) == 1:
        return None

    for i, f in enumerate(files, start=1):
        print(f"{i}. {f.name}")
    n = ask_number("Select a file (0 to quit): ", len(files))
    return files[n - 1] if n else None


def load_tasks(path):
    try:
        wb = load_workbook(path)
    except PermissionError:
        print("Couldn't open the file. Is it open in Excel?")
        return []
    except Exception as e:
        print(f"Couldn't open '{path.name}': {e}")
        return []

    tasks = []
    rows = wb.active.iter_rows(min_row=2, max_col=3, values_only=True)
    for row_num, (name, budget, tracked) in enumerate(rows, start=2):
        budget = to_float(budget)
        if not name or not str(name).strip() or not budget or budget <= 0:
            continue
        tasks.append({
            "row": row_num,
            "task": str(name).strip(),
            "budgeted": budget,
            "tracked": max(0.0, to_float(tracked, 0.0)),
        })

    wb.close()
    return tasks


def save_tasks(path, tasks):
    while True:
        try:
            wb = load_workbook(path)
            sheet = wb.active
            sheet.cell(row=1, column=3).value = "Hours Tracked"
            for t in tasks:
                sheet.cell(row=t["row"], column=3).value = t["tracked"]
            wb.save(path)
            wb.close()
            print("Saved.")
            return True
        except PermissionError:
            print("Couldn't save. Is the file open in Excel?")
        except Exception as e:
            print(f"Couldn't save: {e}")

        if input("ENTER to retry, 'q' to give up: ").strip().lower() == "q":
            print("Changes not saved.")
            return False


def run_timer(task):
    print(f"\nWorking on: {task['task']}")
    input("Press ENTER to start...")

    start = time.monotonic()
    print("Timer running. Press ENTER to stop.")

    while True:
        sys.stdout.write(f"\rElapsed: {clock(time.monotonic() - start)}")
        sys.stdout.flush()
        if msvcrt.kbhit() and msvcrt.getch() == b"\r":
            break
        time.sleep(0.1)

    hours = (time.monotonic() - start) / 3600
    print(f"\nTime worked: {hours:.2f} hours")
    return hours


def adjust_hours(task):
    print(f"\nCurrently tracked: {task['tracked']:.2f} hours")

    while True:
        amount = to_float(input("Hours to add or subtract (0 to cancel): "))

        if amount is None:
            print("Enter a valid number.")
        elif amount == 0:
            return False
        elif task["tracked"] + amount < 0:
            print("Tracked hours can't go below zero.")
        else:
            task["tracked"] += amount
            return True


def task_menu(task):
    """Returns True if the task's hours changed."""
    while True:
        print(f"\n=== {task['task']} ===")
        print(f"Tracked:  {task['tracked']:.2f}")
        print(f"Budgeted: {task['budgeted']:.2f}")
        print("\n1. Start timer\n2. Adjust hours\n3. Back")

        choice = input("> ").strip()
        if choice == "1":
            task["tracked"] += run_timer(task)
            return True
        if choice == "2":
            return adjust_hours(task)
        if choice == "3":
            return False
        print("Pick 1, 2, or 3.")


def select_task(tasks):
    print("\n=== Tasks ===")
    for i, t in enumerate(tasks, start=1):
        print(f"{i}. {t['task']} ({t['tracked']:.2f} / {t['budgeted']:.2f} hrs)")
    print("0. Exit")

    n = ask_number("> ", len(tasks))
    return tasks[n - 1] if n else None


def main():
    path = choose_workbook()
    if path is None:
        print("No file selected.")
        return

    tasks = load_tasks(path)
    if not tasks:
        print("No valid tasks found.")
        return

    unsaved = False
    while True:
        task = select_task(tasks)
        if task is None:
            if unsaved:
                print("\nYou have unsaved changes.")
                save_tasks(path, tasks)
            break

        if task_menu(task):
            print(f"{task['task']}: {task['tracked']:.2f} hours tracked.")
            unsaved = not save_tasks(path, tasks)


if __name__ == "__main__":
    main()